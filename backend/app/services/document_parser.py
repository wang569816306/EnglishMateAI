"""
文档解析服务
支持 Word 和 Excel 文档解析
"""
from io import BytesIO
from typing import BinaryIO


class DocumentParser:
    """文档解析器"""

    @staticmethod
    def parse(file_content: bytes, file_type: str) -> str:
        """
        解析文档内容
        
        Args:
            file_content: 文件字节内容
            file_type: 文件类型 (word/excel)
            
        Returns:
            解析后的文本内容
        """
        if file_type == 'word':
            return DocumentParser._parse_word(file_content)
        elif file_type == 'excel':
            return DocumentParser._parse_excel(file_content)
        else:
            raise ValueError(f"不支持的文件类型: {file_type}")

    @staticmethod
    def _parse_word(file_content: bytes) -> str:
        """
        解析 Word 文档
        
        Args:
            file_content: Word文件字节内容
            
        Returns:
            提取的文本内容
        """
        try:
            # 尝试使用 python-docx 解析 .docx
            from docx import Document
            
            doc = Document(BytesIO(file_content))
            
            # 提取所有段落的文本
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():  # 跳过空段落
                    text_parts.append(para.text.strip())
            
            # 提取表格中的文本
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(' | '.join(row_text))
            
            return '\n\n'.join(text_parts) if text_parts else ''
            
        except ImportError:
            raise ImportError(
                "缺少 python-docx 库，请运行: pip install python-docx"
            )
        except Exception as e:
            raise ValueError(f"Word文档解析失败: {str(e)}")

    @staticmethod
    def _parse_excel(file_content: bytes) -> str:
        """
        解析 Excel 文档
        
        Args:
            file_content: Excel文件字节内容
            
        Returns:
            提取的文本内容
        """
        try:
            # 尝试使用 openpyxl 解析 .xlsx
            from openpyxl import load_workbook
            
            wb = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
            
            text_parts = []
            
            # 遍历所有工作表
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                if ws.max_row == 0:
                    continue
                
                text_parts.append(f"【工作表: {sheet_name}】")
                
                # 提取单元格内容
                for row in ws.iter_rows(values_only=True):
                    # 过滤空行
                    row_values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if row_values:
                        text_parts.append('\t'.join(row_values))
                
                text_parts.append("")  # 工作表之间空一行
            
            wb.close()
            return '\n'.join(text_parts) if text_parts else ''
            
        except ImportError:
            raise ImportError(
                "缺少 openpyxl 库，请运行: pip install openpyxl"
            )
        except Exception as e:
            # 尝试使用 xlrd 解析 .xls
            try:
                import xlrd
                
                workbook = xlrd.open_workbook(file_contents=file_content)
                text_parts = []
                
                for sheet_name in workbook.sheet_names():
                    ws = workbook.sheet_by_name(sheet_name)
                    
                    if ws.nrows == 0:
                        continue
                    
                    text_parts.append(f"【工作表: {sheet_name}】")
                    
                    for row_idx in range(ws.nrows):
                        row_values = []
                        for col_idx in range(ws.ncols):
                            cell_value = ws.cell_value(row_idx, col_idx)
                            if cell_value is not None and str(cell_value).strip():
                                row_values.append(str(cell_value).strip())
                        
                        if row_values:
                            text_parts.append('\t'.join(row_values))
                    
                    text_parts.append("")
                
                return '\n'.join(text_parts) if text_parts else ''
                
            except ImportError:
                raise ImportError(
                    "缺少 openpyxl 或 xlrd 库，请运行: pip install openpyxl xlrd"
                )
            except Exception as e2:
                raise ValueError(f"Excel文档解析失败: {str(e)}")
